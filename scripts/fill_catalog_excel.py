#!/usr/bin/env python3
"""Extract products from GTC HTML pages and fill Meta catalog Excel template."""

import json
import re
import subprocess
import time
import urllib.request
from pathlib import Path
from typing import Optional
from urllib.parse import urlsplit, urlunsplit

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BASE_URL = "https://www.gayathrithreadcouture.in"
BRAND = "Gayathri Thread Couture"
GOOGLE_CATEGORY = "Apparel & Accessories > Jewelry"
FB_CATEGORY = "Clothing & Accessories > Jewelry"
PRICE_FALLBACK = 0
FETCH_LIVE_PRICES = False

PAGES = [
    {
        "file": "html/bands.html",
        "prefix": "band",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/band-{p['id']}",
        "page_url": f"{BASE_URL}/html/bands.html",
    },
    {
        "file": "html/bangles.html",
        "prefix": "bangle",
        "variable": "silkThreadProducts",
        "link": lambda p: (
            f"https://gayathrithreadcouture.myinstamojo.com/product/bangle-{p.get('storeId', p['id'])}"
            + (f"?Colour={p['colour']}" if p.get("colour") else "")
        ),
        "page_url": f"{BASE_URL}/html/bangles.html",
    },
    {
        "file": "html/bangles.html",
        "prefix": "bridal-bangle",
        "variable": "bridalProducts",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/bangle-{p.get('storeId', p['id'])}",
        "page_url": f"{BASE_URL}/html/bangles.html",
    },
    {
        "file": "html/bracelets.html",
        "prefix": "bracelet",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/bracelet-{p.get('storeId', p['id'])}",
        "page_url": f"{BASE_URL}/html/bracelets.html",
    },
    {
        "file": "html/centerclips.html",
        "prefix": "center",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/center-{p['id']}",
        "page_url": f"{BASE_URL}/html/centerclips.html",
    },
    {
        "file": "html/chains.html",
        "prefix": "chain",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/chain-{p['id']}",
        "page_url": f"{BASE_URL}/html/chains.html",
    },
    {
        "file": "html/clips.html",
        "prefix": "clip",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/clip-{p['id']}",
        "page_url": f"{BASE_URL}/html/clips.html",
    },
    {
        "file": "html/pins.html",
        "prefix": "pins",
        "variable": "products",
        "link": lambda p: f"https://gayathrithreadcouture.myinstamojo.com/product/pins-{p['id']}",
        "page_url": f"{BASE_URL}/html/pins.html",
    },
]

EXTRACT_JS = r"""
const fs = require('fs');
const html = fs.readFileSync(process.argv[1], 'utf8');
const variable = process.argv[2] || 'products';
const marker = `const ${variable} = [`;
const start = html.indexOf(marker);
if (start === -1) { console.log('[]'); process.exit(0); }
let i = start + marker.length - 1;
let depth = 0;
let inStr = false;
let strCh = '';
let escaped = false;
for (; i < html.length; i++) {
  const ch = html[i];
  if (inStr) {
    if (escaped) { escaped = false; continue; }
    if (ch === '\\') { escaped = true; continue; }
    if (ch === strCh) inStr = false;
    continue;
  }
  if (ch === '"' || ch === "'") { inStr = true; strCh = ch; continue; }
  if (ch === '[') depth++;
  if (ch === ']') {
    depth--;
    if (depth === 0) { i++; break; }
  }
}
    const arr = html.slice(start + marker.length - 1, i);
try {
  const products = eval(arr);
  console.log(JSON.stringify(products));
} catch (e) {
  console.error(e.message);
  process.exit(1);
}
"""


def extract_products(html_path: Path, variable: str = "products") -> list:
    result = subprocess.run(
        ["node", "-e", EXTRACT_JS, str(html_path), variable],
        capture_output=True,
        text=True,
        check=True,
    )
    return json.loads(result.stdout or "[]")


def to_image_url(image: str) -> str:
    if image.startswith("http"):
        return canonical_url(image)
    return canonical_url(BASE_URL + "/" + image.lstrip("./").replace("../", ""))


def canonical_url(url: str) -> str:
    parts = urlsplit(url.strip())
    hostname = (parts.hostname or "").lower().removeprefix("www.")
    netloc = hostname
    if parts.port:
        netloc += f":{parts.port}"
    return urlunsplit((parts.scheme.lower(), netloc, parts.path, parts.query, parts.fragment))


def fetch_instamojo_price(url: str, cache: dict) -> str:
    if not FETCH_LIVE_PRICES:
        return ""
    if url in cache:
        return cache[url]
    try:
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (compatible; GTC-CatalogBot/1.0)"},
        )
        with urllib.request.urlopen(req, timeout=5) as resp:
            html = resp.read().decode("utf-8", errors="replace")
        match = re.search(
            r'property="product:price:amount"\s+content="([^"]+)"', html
        )
        if match:
            amount = match.group(1).strip()
            price = f"{float(amount):.2f} INR"
        else:
            price = ""
    except Exception:
        price = ""
    cache[url] = price
    time.sleep(0.05)
    return price


def build_description(product: dict, page_label: str) -> str:
    parts = []
    if product.get("description"):
        parts.append(product["description"])
    if product.get("materials"):
        parts.append(f"Materials: {product['materials']}")
    if product.get("occasions"):
        parts.append(f"Perfect for: {product['occasions']}")
    if not parts:
        parts.append(f"Handcrafted {page_label} from {BRAND}.")
    return " ".join(parts)


def infer_color(name: str) -> Optional[str]:
    colors = [
        "black", "white", "pink", "red", "blue", "green", "gold", "golden", "silver",
        "purple", "yellow", "orange", "maroon", "burgundy", "navy", "teal", "turquoise",
        "lavender", "coral", "cream", "ivory", "ruby", "emerald", "amber", "peach",
        "mint", "rose", "magenta", "cyan", "khaki", "olive", "slate", "blush", "wine",
        "salmon", "champagne", "plum", "forest", "steel", "dusty rose", "periwinkle",
        "garnet", "celadon", "berry", "cornflower", "tangerine", "taupe", "raspberry",
        "mustard", "peacock", "aqua", "lemon", "rainbow", "seafoam", "royal blue",
    ]
    lower = name.lower()
    for color in colors:
        if color in lower:
            return color.title() if " " not in color else color
    return None


def product_row(product: dict, config: dict, price_cache: dict) -> dict:
    prefix = config["prefix"]
    product_id = product.get("id")
    content_id = f"{prefix}-{product_id}"
    name = product.get("name", f"{prefix.title()} {product_id}")
    image = to_image_url(product.get("image", ""))
    description = build_description(product, config["prefix"].replace("-", " "))
    color = product.get("colour") or infer_color(name)
    material = product.get("materials")
    category_tag = product.get("category", config["prefix"].replace("-", " ").title())
    link = config["link"](product)
    price = PRICE_FALLBACK

    row = {
        "id": content_id,
        "title": name[:200],
        "description": description[:9999],
        "availability": "in stock",
        "condition": "new",
        "price": price,
        "link": link,
        "image_link": image,
        "brand": BRAND,
        "google_product_category": GOOGLE_CATEGORY,
        "fb_product_category": FB_CATEGORY,
        "gender": "female",
        "age_group": "adult",
        "product_tags[0]": category_tag,
        "product_tags[1]": config["prefix"],
    }
    if material:
        row["material"] = material[:200]
    if color:
        row["color"] = str(color)[:200]
    return row


def main():
    template = ROOT / "catalog_products_filled.xlsx"
    output = template

    wb = openpyxl.load_workbook(template)
    ws = wb.active

    headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    header_index = {h: i + 1 for i, h in enumerate(headers) if h}

    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)

    all_rows = []
    price_cache = {}
    for config in PAGES:
        html_path = ROOT / config["file"]
        products = extract_products(html_path, config.get("variable", "products"))
        print(f"{config['file']} ({config.get('variable', 'products')}): {len(products)} products")
        for product in products:
            all_rows.append(product_row(product, config, price_cache))

    print(f"Total products: {len(all_rows)}")
    priced = sum(1 for r in all_rows if r.get("price"))
    print(f"Prices fetched: {priced}/{len(all_rows)}")

    unique_rows = []
    seen_images = set()
    for row_data in all_rows:
        image = canonical_url(row_data.get("image_link", ""))
        row_data["image_link"] = image
        if image and image in seen_images:
            continue
        if image:
            seen_images.add(image)
        unique_rows.append(row_data)

    print(f"Unique image URLs: {len(unique_rows)}")
    for row_data in unique_rows:
        row_num = ws.max_row + 1
        for field, col in header_index.items():
            if field in row_data and row_data[field] is not None:
                ws.cell(row_num, col, row_data[field])

    wb.save(output)
    print(f"Saved: {output}")


if __name__ == "__main__":
    main()
